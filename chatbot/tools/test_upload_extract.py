"""
test_upload_extract.py — unit tests for tools/upload_extract.py.

No LLM, no CrewAI, no FAISS, no network — this only exercises the plain
extract_text() / build_uploaded_document_block() functions against real,
locally-generated PDF/DOCX fixtures. Safe to run anywhere, including this
sandbox (unlike the agents/test_*.py scripts, which need a live Ollama or
Anthropic endpoint).

Fixture generation needs `reportlab` to synthesize a real text-bearing PDF
(`pip install reportlab --break-system-packages`) — this is a TEST-ONLY
dependency, deliberately not added to chatbot/requirements.txt, since
upload_extract.py itself never writes PDFs, only reads them. python-docx
is already a runtime dependency (see requirements.txt) and is reused here
to build the DOCX fixture too.

Run:
    python chatbot/tools/test_upload_extract.py
"""

import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import upload_extract as ue  # noqa: E402

_FAILURES = []


def _check(label: str, condition: bool, detail: str = "") -> None:
    status = "PASS" if condition else "FAIL"
    print(f"[{status}] {label}" + (f" — {detail}" if detail and not condition else ""))
    if not condition:
        _FAILURES.append(label)


def _make_pdf(path: str, lines: list[str]) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(path, pagesize=(612, 792))
    c.setFont("Helvetica", 12)
    y = 720
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.showPage()
    c.save()


def _make_docx(path: str, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> None:
    import docx

    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row_vals in table_rows:
            row = table.add_row()
            for cell, val in zip(row.cells, row_vals):
                cell.text = val
    doc.save(path)


def test_pdf_extraction(tmpdir: str) -> None:
    path = os.path.join(tmpdir, "syllabus.pdf")
    _make_pdf(
        path,
        [
            "AI Curriculum Draft — Department of Computer Science",
            "Course: Introduction to Machine Learning",
            "Required: Linear Algebra, Probability, Python Programming",
        ],
    )
    text = ue.extract_text(path)
    _check("pdf: non-empty text extracted", bool(text.strip()))
    _check(
        "pdf: contains expected content",
        "Introduction to Machine Learning" in text,
        detail=text[:200],
    )


def test_docx_extraction(tmpdir: str) -> None:
    path = os.path.join(tmpdir, "syllabus.docx")
    _make_docx(
        path,
        paragraphs=[
            "AI Curriculum Draft — Department of Computer Science",
            "This program covers applied machine learning and ethics.",
        ],
        table_rows=[
            ["Course", "Type"],
            ["Intro to ML", "Required"],
            ["NLP Seminar", "Elective"],
        ],
    )
    text = ue.extract_text(path)
    _check("docx: non-empty text extracted", bool(text.strip()))
    _check(
        "docx: paragraph content present",
        "applied machine learning and ethics" in text,
        detail=text[:200],
    )
    _check(
        "docx: table content present (cells joined with ' | ')",
        "Intro to ML | Required" in text,
        detail=text,
    )


def test_truncation(tmpdir: str) -> None:
    path = os.path.join(tmpdir, "long.docx")
    # One paragraph per line, enough repeated lines to comfortably exceed
    # MAX_EXTRACTED_CHARS (12,000).
    long_line = "This is a long curriculum description line used to force truncation. "
    paragraphs = [long_line] * 400  # ~400 * 72 chars ≈ 28,800 chars
    _make_docx(path, paragraphs=paragraphs)
    text = ue.extract_text(path)
    _check(
        "truncation: output capped near MAX_EXTRACTED_CHARS",
        len(text) <= ue.MAX_EXTRACTED_CHARS + 200,
        detail=f"len={len(text)}",
    )
    _check("truncation: note appended", "TRUNCATED" in text)


def test_unsupported_type(tmpdir: str) -> None:
    path = os.path.join(tmpdir, "notes.txt")
    with open(path, "w") as f:
        f.write("plain text file, not pdf/docx")
    try:
        ue.extract_text(path)
        _check("unsupported type: raises UnsupportedFileTypeError", False, "no exception raised")
    except ue.UnsupportedFileTypeError:
        _check("unsupported type: raises UnsupportedFileTypeError", True)
    except Exception as exc:  # noqa: BLE001
        _check(
            "unsupported type: raises UnsupportedFileTypeError",
            False,
            f"raised {type(exc).__name__} instead",
        )


def test_missing_file() -> None:
    try:
        ue.extract_text("/tmp/this_file_does_not_exist_12345.pdf")
        _check("missing file: raises FileNotFoundError", False, "no exception raised")
    except FileNotFoundError:
        _check("missing file: raises FileNotFoundError", True)
    except Exception as exc:  # noqa: BLE001
        _check("missing file: raises FileNotFoundError", False, f"raised {type(exc).__name__} instead")


def test_corrupted_pdf(tmpdir: str) -> None:
    path = os.path.join(tmpdir, "corrupted.pdf")
    with open(path, "wb") as f:
        f.write(b"%PDF-1.4\nthis is not a real pdf body, just garbage bytes\n%%EOF")
    try:
        ue.extract_text(path)
        _check("corrupted pdf: raises ExtractionError", False, "no exception raised")
    except ue.ExtractionError:
        _check("corrupted pdf: raises ExtractionError", True)
    except Exception as exc:  # noqa: BLE001
        _check("corrupted pdf: raises ExtractionError", False, f"raised {type(exc).__name__} instead")


def test_empty_pdf_no_text(tmpdir: str) -> None:
    from reportlab.pdfgen import canvas

    path = os.path.join(tmpdir, "blank.pdf")
    c = canvas.Canvas(path, pagesize=(612, 792))
    c.showPage()  # blank page, no text drawn
    c.save()
    try:
        ue.extract_text(path)
        _check("blank pdf (no text): raises ExtractionError", False, "no exception raised")
    except ue.ExtractionError:
        _check("blank pdf (no text): raises ExtractionError", True)
    except Exception as exc:  # noqa: BLE001
        _check("blank pdf (no text): raises ExtractionError", False, f"raised {type(exc).__name__} instead")


def test_build_uploaded_document_block() -> None:
    block = ue.build_uploaded_document_block("syllabus.pdf", "Intro to ML, taught in fall.")
    _check(
        "block: contains marker string agents pattern-match on",
        "ATTACHED UPLOADED CURRICULUM DOCUMENT" in block,
    )
    _check("block: filename present", "syllabus.pdf" in block)
    _check(
        "block: provenance/not-verified framing present",
        "NOT independently verified" in block,
    )
    _check("block: extracted text included verbatim", "Intro to ML, taught in fall." in block)


def main() -> int:
    with tempfile.TemporaryDirectory() as tmpdir:
        test_pdf_extraction(tmpdir)
        test_docx_extraction(tmpdir)
        test_truncation(tmpdir)
        test_unsupported_type(tmpdir)
        test_missing_file()
        test_corrupted_pdf(tmpdir)
        test_empty_pdf_no_text(tmpdir)
        test_build_uploaded_document_block()

    print()
    if _FAILURES:
        print(f"{len(_FAILURES)} FAILURE(S): {_FAILURES}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
