"""
upload_extract.py — text extraction for professor-uploaded curriculum files.

Backs the "upload an unpublished/not-yet-official curriculum file" feature
(added 2026-06-23). Unlike every other corpus in this project
(chatbot/data/program_and_curriculum/, the news/industry-reports corpus),
content extracted here is NEVER written to a FAISS index or any persistent
store — per the 2026-06-23 scoping decision, this is a one-off,
in-conversation analysis feature only, not a new searchable corpus.
extract_text() is called once per uploaded file, directly from
chatbot/app.py's on_message handler, and the resulting text is folded
straight into that single query's Task description (see
build_uploaded_document_block()). Nothing here touches embeddings, FAISS,
or any of the *_rag_tool.py modules.

This is a deliberate departure from this project's "every corpus entry is
human-verified against an official published page before being added"
discipline (see data/program_and_curriculum/README.md) — an unpublished/
not-yet-official file BY DEFINITION has no official page to verify against.
build_uploaded_document_block() exists specifically to keep that distinction
visible to the agent: the wrapped block states explicitly that this content
is professor-provided and not independently verified, so
agents/university_programs.py's backstory (see its "HANDLING AN ATTACHED
UPLOADED CURRICULUM DOCUMENT" section) can cite it correctly instead of
treating it like a verified program_rag_tool/web_search_tool hit.

First version (2026-06-23) supports PDF and DOCX only — the two formats a
course syllabus/curriculum draft is most likely to arrive in (per explicit
user scoping decision). Anything else raises UnsupportedFileTypeError with
a clear message; extend _EXTRACTORS to add a format later (e.g. .pptx,
.txt).

Run standalone to smoke-test against a real file:
    python chatbot/tools/upload_extract.py /path/to/some.pdf
"""

import os
import sys

# ~3,000 tokens — generous for a syllabus/course-outline document, but
# bounded so one huge upload can't blow out the agent's context window.
# See CLAUDE.md's "Ollama Context Window" section for why this project is
# conservative about unbounded prompt length.
MAX_EXTRACTED_CHARS = 12_000

_TRUNCATION_NOTE_TEMPLATE = (
    "\n\n[... TRUNCATED — original extracted text was {orig} characters, "
    "showing the first {cap} ...]"
)


class UnsupportedFileTypeError(ValueError):
    """Raised when the uploaded file's extension isn't PDF or DOCX (v1 scope)."""


class ExtractionError(RuntimeError):
    """Raised when a supported file type fails to parse (corrupted, encrypted,
    scanned/image-only with no extractable text, etc.)."""


def _extract_pdf(file_path: str) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(file_path)
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF {file_path!r}: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        raise ExtractionError(
            f"PDF {file_path!r} is password-protected — cannot extract text."
        )

    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            # Don't let one bad page kill the whole extraction — note it
            # inline and keep going, consistent with this project's
            # "degrade gracefully, flag the gap" convention (see
            # web_search_tool.py / news_tool.py).
            pages.append(f"[page {i + 1}: extraction failed — {exc}]")
    return "\n\n".join(pages).strip()


def _extract_docx(file_path: str) -> str:
    import docx

    try:
        doc = docx.Document(file_path)
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX {file_path!r}: {exc}") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables often hold the actual course list / weekly schedule in real
    # syllabus documents — extract their cell text too, row by row, so it
    # isn't silently dropped.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def extract_text(file_path: str) -> str:
    """
    Extract plain text from an uploaded PDF or DOCX file.

    Raises:
      FileNotFoundError        — file_path doesn't exist.
      UnsupportedFileTypeError — extension isn't .pdf or .docx (v1 scope).
      ExtractionError          — matched a supported extension but failed
                                  to parse (corrupted, encrypted, or zero
                                  extractable text — e.g. a scanned/
                                  image-only PDF this version can't OCR).

    Output is capped at MAX_EXTRACTED_CHARS characters with a trailing
    truncation note if the original text was longer — see module docstring.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(file_path)

    ext = os.path.splitext(file_path)[1].lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type {ext!r}. This version only accepts PDF "
            f"and DOCX curriculum uploads (got: {os.path.basename(file_path)})."
        )

    text = extractor(file_path)
    if not text.strip():
        raise ExtractionError(
            f"No extractable text found in {os.path.basename(file_path)!r} "
            "— it may be a scanned/image-only document, which this version "
            "cannot OCR."
        )

    if len(text) > MAX_EXTRACTED_CHARS:
        orig_len = len(text)
        text = text[:MAX_EXTRACTED_CHARS] + _TRUNCATION_NOTE_TEMPLATE.format(
            orig=orig_len, cap=MAX_EXTRACTED_CHARS
        )

    return text


def build_uploaded_document_block(filename: str, text: str) -> str:
    """
    Wrap extracted text in a clearly-delimited block for injection directly
    into a CrewAI Task description (see chatbot/app.py's on_message).

    The wording here is load-bearing, not decorative: agents/orchestrator.py
    and agents/university_programs.py both instruct the LLM to pattern-match
    on the literal "ATTACHED UPLOADED CURRICULUM DOCUMENT" marker string in
    their backstories to decide how to route and cite this content. Keep
    that marker text in sync across all three files if you ever change it.
    """
    return (
        "===== ATTACHED UPLOADED CURRICULUM DOCUMENT =====\n"
        f"Filename: {filename}\n"
        "IMPORTANT PROVENANCE NOTE: this document was uploaded directly by "
        "the professor/user in this conversation. Unlike the local "
        "program_and_curriculum corpus or live web search results, it is "
        "NOT independently verified against any official published page — "
        "by definition (it's unpublished / not-yet-official). Treat its "
        "content as an authoritative description of what THIS professor's "
        "own program/course currently contains, but always cite it "
        "explicitly as \"from the uploaded document\" or \"as provided by "
        "the professor\" — never as a verified peer-institution source.\n"
        "--- Extracted text begins ---\n"
        f"{text}\n"
        "--- Extracted text ends ---\n"
        "===== END ATTACHED DOCUMENT ====="
    )


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python chatbot/tools/upload_extract.py <file.pdf|file.docx>")
        sys.exit(1)
    path = sys.argv[1]
    try:
        extracted = extract_text(path)
    except (UnsupportedFileTypeError, ExtractionError, FileNotFoundError) as exc:
        print(f"❌ {type(exc).__name__}: {exc}")
        sys.exit(1)
    print(f"Extracted {len(extracted)} chars from {path}\n")
    print(build_uploaded_document_block(os.path.basename(path), extracted))
